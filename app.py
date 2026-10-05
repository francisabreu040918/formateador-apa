import io
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai

# 1. Configuración de la interfaz
st.set_page_config(
    page_title="Formateador APA 7 Automático",
    page_icon="🎓",
    layout="centered"
)

st.title("🎓 Generador & Formateador APA 7")
st.markdown(
    "Pega el texto de tu trabajo o tu lista de fuentes desordenadas. "
    "La herramienta organizará las citas y referencias en **APA 7ma edición** "
    "y te generará el documento **Word (.docx)** listo para entregar."
)

# 2. Obtener la clave de API desde los secretos de Streamlit
api_key = st.secrets.get("GEMINI_API_KEY", "")

# 3. Formulario principal
tipo_trabajo = st.selectbox(
    "¿Qué deseas procesar hoy?",
    [
        "Referencias bibliográficas (Lista final en orden alfabético)",
        "Párrafo/Texto con citas parentéticas y narrativas",
        "Ensayo o resumen completo"
    ]
)

texto_usuario = st.text_area(
    "Pega tu texto, enlaces o fuentes aquí:",
    height=200,
    placeholder="Ejemplo:\n- Sampieri metodología 2014 McGrawHill\n- https://scielo.org/articulo-ejemplo\n- Vygotsky teoría sociocultural..."
)

# 4. Función para crear el archivo Word con las reglas exactas de APA 7ma edición
def generar_word_apa(texto_procesado, titulo="Trabajo Académico en Formato APA"):
    doc = Document()

    # Márgenes oficiales APA 7: 1 pulgada (2.54 cm) en todos los lados
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Tipografía oficial: Times New Roman 12 pt, interlineado doble
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.space_after = Pt(0)

    # Título centrado en negrita (Nivel 1 APA)
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = p_titulo.add_run(titulo)
    run_titulo.bold = True

    # Recorrer las líneas devueltas por la IA
    lineas = texto_procesado.split("\n")
    es_seccion_referencias = False

    for linea in lineas:
        linea_limpia = linea.strip()
        if not linea_limpia:
            continue

        # Si detecta el encabezado de referencias
        if "referencias" in linea_limpia.lower() and len(linea_limpia) < 25:
            es_seccion_referencias = True
            p_ref = doc.add_paragraph()
            p_ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_ref = p_ref.add_run("Referencias")
            run_ref.bold = True
            continue

        p = doc.add_paragraph()

        if es_seccion_referencias:
            # Sangría francesa reglamentaria en APA: 0.5 pulgadas (1.27 cm)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(linea_limpia)
        else:
            # Sangría normal de primera línea en APA
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(linea_limpia)

    # Guardar en memoria para descarga sin escribir en disco
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# 5. Botón para procesar con la IA
if st.button("⚡ Procesar en Formato APA 7", type="primary"):
    if not api_key:
        st.error("⚠️ Falta configurar la GEMINI_API_KEY en los Secrets de Streamlit.")
    elif not texto_usuario.strip():
        st.warning("⚠️ Debes pegar algún texto o fuente en el cuadro superior.")
    else:
        with st.spinner("Organizando fuentes y aplicando normas APA 7ma edición..."):
            try:
                client = genai.Client(api_key=api_key)

                instrucciones = (
                    "Eres un experto metodólogo universitario y revisor de estilo en normas APA 7ma edición.\n"
                    f"El usuario seleccionó: {tipo_trabajo}.\n\n"
                    "Instrucciones estrictas:\n"
                    "1. Transforma el contenido al formato oficial APA 7ma edición.\n"
                    "2. Si son referencias, ordénalas alfabéticamente por el apellido del autor. "
                    "Asegura la estructura: Apellido, Inicial. (Año). Título en cursiva. Editorial/Revista, DOI o URL.\n"
                    "3. Si es texto, verifica y corrige las citas parentéticas (Apellido, Año).\n"
                    "4. No agregues introducciones, saludos ni comentarios. Solo entrega el contenido final listo para el documento, "
                    "encabezado por 'Referencias' si corresponde.\n\n"
                    "Contenido a procesar:\n"
                    f"{texto_usuario}"
                )

                response = client.models.generate_content(
                    model="gemini-3.8-flash",
                    contents=instrucciones
                )

                st.session_state["resultado_apa"] = response.text
                st.success("¡Contenido formateado exitosamente!")

            except Exception as e:
                st.error(f"Error al procesar: {str(e)}")

# 6. Mostrar el resultado y el bloqueo de cobro
if "resultado_apa" in st.session_state:
    st.markdown("---")
    st.subheader("📄 Vista Previa del Resultado")
    st.text_area("Contenido generado:", value=st.session_state["resultado_apa"], height=180)

    st.markdown("---")
    st.subheader("📥 Descargar Documento Word (.docx)")
    st.info("El archivo Word incluye los márgenes oficiales (2.54 cm), fuente Times New Roman 12, interlineado doble y la sangría francesa ya configurada.")

    # Código de desbloqueo
    CODIGO_VALIDO = "APA2026"

    col1, col2 = st.columns([2, 1])
    with col1:
        codigo_ingresado = st.text_input("Introduce tu código de acceso para desbloquear la descarga:", type="password")
    with col2:
        st.markdown("<br>", unsafe_allow_html=True)
        numero_whatsapp = "18090000000" 
        mensaje_ws = "Hola! Quiero mi código de acceso para descargar mi trabajo en formato APA."
        url_whatsapp = f"https://wa.me/{numero_whatsapp}?text={mensaje_ws.replace(' ', '%20')}"
        st.markdown(f"[📲 Solicitar código por WhatsApp]({url_whatsapp})")

    if codigo_ingresado == CODIGO_VALIDO:
        st.success("¡Código correcto! Ya puedes descargar tu archivo:")
        archivo_word = generar_word_apa(st.session_state["resultado_apa"])
        st.download_button(
            label="⬇️ Descargar archivo Word (.docx)",
            data=archivo_word,
            file_name="Trabajo_Formato_APA7.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    elif codigo_ingresado != "":
        st.error("Código incorrecto. Solicita tu código vía WhatsApp.")
